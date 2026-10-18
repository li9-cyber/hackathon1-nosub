"""Synthetic CVs only: no user data or provider calls in the default suite."""
from io import BytesIO
import json
import pymupdf
from docx import Document
from career_twin.profile.schema import Extraction

NAMES = ['Python', 'SQL', 'MS Excel', 'Tableau', 'Power BI', 'pandas', 'Java', 'R',
         'Rust', 'Git', 'Docker', 'Bloomberg Terminal', 'Figma', 'VBA', 'SAP', 'Snowflake', 'dbt', 'CustomTool-X']
WORK = 'EXPERIENCE\nExample Bank | Finance Intern | Jun-Aug 2025\nUsed Python and SQL to reconcile accounts.\nBuilt charts with Python and MS Excel.'
SKILLS = 'SKILLS\n' + ', '.join(NAMES)
EDUCATION = 'EDUCATION\nExample University | BSc Finance | 2022-2026'
CERTIFICATE = 'CERTIFICATION\nCFA Level I passed | 2025'
TEXT = '\n\n'.join([WORK, SKILLS, EDUCATION, CERTIFICATE])
SUPPLEMENT = 'Supplement 😀: Used Docker and Python to automate reports.\nEarned the AWS Cloud Practitioner certificate in 2026.'


def supplement_result(text):
    return Extraction.model_validate({'entries': [dict(section='experience', title='Supplement', source_quote=text,
        activities=['Used Docker and Python to automate reports.'],
        fields={'資格': 'AWS Cloud Practitioner', '日期': '2026'},
        skills=[dict(name='Docker', quotes=['Used Docker and Python to automate reports.']),
                dict(name='Python', quotes=['Used Docker and Python to automate reports.'])])]})


def model_result(text):
    # Use the exact parser output, including PDF line wraps / DOCX table tabs.
    def section(start, end=None):
        return text[text.index(start):text.index(end) if end else len(text)].rstrip('\n')
    work, skills = section('EXPERIENCE', 'SKILLS'), section('SKILLS', 'EDUCATION')
    return Extraction.model_validate({'entries': [
        dict(section='experience', title='Finance Intern', source_quote=work,
             organization='Example Bank', role='Finance Intern', date_range='Jun-Aug 2025', experience_type='internship',
             activities=['Used Python and SQL to reconcile accounts.', 'Built charts with Python and MS Excel.'],
             skills=[dict(name='Python', quotes=['Used Python and SQL to reconcile accounts.', 'Built charts with Python and MS Excel.']),
                     dict(name='SQL', quotes=['Used Python and SQL to reconcile accounts.']),
                     dict(name='MS Excel', quotes=['Built charts with Python and MS Excel.'])]),
        dict(section='skill', title='Skills', source_quote=skills, skills=[dict(name=name, quotes=[name]) for name in NAMES]),
        dict(section='education', title='BSc Finance', source_quote=section('EDUCATION', 'CERTIFICATION'),
             fields={'學校': 'Example University', '學位': 'BSc Finance', '日期': '2022-2026'}),
        dict(section='credential', title='CFA Level I', source_quote=section('CERTIFICATION'), fields={'證照': 'CFA Level I passed', '日期': '2025'})
    ]})


def document_bytes(extension, pages=1):
    if extension == 'txt':
        return TEXT.encode('utf-8')
    if extension == 'pdf':
        with pymupdf.open() as document:
            chunks = [TEXT] if pages == 1 else [WORK, SKILLS, EDUCATION + '\n\n' + CERTIFICATE]
            for text in chunks:
                page = document.new_page()
                assert page.insert_textbox(pymupdf.Rect(40, 40, 550, 800), text, fontsize=11) >= 0
            return document.tobytes()
    document = Document()
    for line in WORK.splitlines():
        document.add_paragraph(line)
    document.add_paragraph('SKILLS')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = ', '.join(NAMES[:9])
    table.cell(0, 1).text = ', '.join(NAMES[9:])
    for line in (EDUCATION + '\n' + CERTIFICATE).splitlines():
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def mock_response(request):
    payload = json.loads(request.content)
    text = json.loads(payload['messages'][1]['content'])['cv_text']
    return {'choices': [{'finish_reason': 'stop', 'message': {'content': model_result(text).model_dump_json()}}]}
